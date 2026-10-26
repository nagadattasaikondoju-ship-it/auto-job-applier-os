"""ATS resume scanner: text extraction, deterministic scoring and truthful rewrite checks.

Scoring is rule-based so it is fast, explainable and works without an AI key. AI is only
used for bullet rewrites, and every rewrite is checked so it can't add facts.
"""

import io
import re
from typing import Any

from packages.core.seed.taxonomies import ALL_SKILLS

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

ACTION_VERBS = {
    "achieved", "analysed", "analyzed", "architected", "automated", "built", "championed",
    "collaborated", "configured", "created", "cut", "debugged", "decreased", "delivered",
    "deployed", "designed", "developed", "drove", "enabled", "engineered", "established",
    "evaluated", "expanded", "generated", "grew", "handled", "identified", "implemented",
    "improved", "increased", "integrated", "introduced", "launched", "led", "maintained",
    "managed", "mentored", "migrated", "modernized", "negotiated", "optimised", "optimized",
    "orchestrated", "organized", "owned", "partnered", "planned", "presented", "produced",
    "reduced", "redesigned", "refactored", "resolved", "restructured", "revamped", "scaled",
    "shipped", "simplified", "spearheaded", "standardized", "streamlined", "supported",
    "tested", "trained", "transformed", "tripled", "doubled", "upgraded", "wrote",
}
WEAK_PHRASES = [
    "responsible for", "worked on", "helped with", "helped to", "involved in", "duties included",
    "tasked with", "participated in", "assisted with", "various", "etc",
]
CLICHES = [
    "team player", "hard working", "hardworking", "go-getter", "detail oriented", "detail-oriented",
    "self-motivated", "results-driven", "think outside the box", "synergy", "dynamic",
]
SECTION_PATTERNS = {
    "summary": r"\b(summary|profile|objective|about me)\b",
    "experience": r"\b(experience|employment|work history|internships?)\b",
    "education": r"\b(education|academics?|qualifications?)\b",
    "skills": r"\b(skills|technical skills|tech stack|technologies|tools)\b",
    "projects": r"\b(projects?)\b",
    "certifications": r"\b(certifications?|certificates?|courses)\b",
}
_STOPWORDS = set(
    ["a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "has", "have", "in", "is", "it", "its", "of", "on", "or", "our", "the", "to", "we", "will", "with", "you", "your", "this", "that", "their", "they", "who", "what", "when", "where", "which", "while", "into", "across", "about", "using", "use", "used", "able", "work", "working", "team", "role", "job", "candidate", "experience", "years", "year", "strong", "good", "great", "ability", "skills", "including", "include", "etc", "must", "should", "plus", "nice", "preferred", "required", "requirements", "responsibilities"]
)


class UnsupportedFileError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or text resume."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedFileError("File is larger than 5 MB.")
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileError("Upload a PDF, DOCX or .txt file.")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 100:
        raise UnsupportedFileError(
            "We couldn't read text from this file. If it's a scanned PDF, upload a DOCX or text version."
        )
    return text


def _lines(text: str) -> list[str]:
    return [ln.strip() for ln in text.splitlines() if ln.strip()]


def _bullets(text: str) -> list[str]:
    """Lines that look like achievement bullets."""
    out = []
    for ln in _lines(text):
        stripped = re.sub(r"^[\-•\*●▪–>•]+\s*", "", ln)
        words = stripped.split()
        looks_like_bullet = stripped != ln or (
            len(words) >= 6 and not stripped.isupper() and not stripped.endswith(":")
        )
        if looks_like_bullet and len(words) >= 4:
            out.append(stripped)
    return out


def has_metric(s: str) -> bool:
    return bool(re.search(r"\d+\s*(%|x|k|m|cr|lpa|l\b|ms|hrs?|users|customers|clients)|[₹$]\s?\d|\d{2,}", s, re.I))


# Short skill names that are real skills in job posts even though they're plain capitals
_TECH_ACRONYMS = {
    "ai", "ml", "nlp", "llm", "llms", "ui", "ux", "qa", "bi", "etl", "elt", "api", "apis", "rest", "saas", "seo",
    "sem", "crm", "erp", "sap", "gcp", "azure", "mlops", "devops", "sre", "tdd", "oop", "dsa", "sql", "nosql",
}

# Skills that are also everyday English words only count when capitalised ("Go", not "go live")
_CASE_SENSITIVE_SKILLS = {"Go", "Express", "Swift", "Excel", "Spark", "Git", "Vite", "Android", "Linux"}

# Levels, role abbreviations, degrees and web domains look like tech tokens but aren't skills
_NOT_A_SKILL = re.compile(
    r"^(I{1,3}|IV|VI{0,3}|SDE\d?|SWE|SDET|(?!(AI|ML|UI|UX|QA|BI|R)$)[A-Z]{1,2}|[A-Z]\d"  # levels, role abbreviations, initials ("RV")
    r"|[BM]\.?(Tech|E|Sc|Com|A|S)|MBA|BCA|MCA|PhD|CGPA|GPA|SGPA"  # degrees and grades
    r"|ROI|KPIs?|OKRs?|SLAs?|P&L|CTC|LPA|USD|INR"  # business terms, not skills
    r"|[A-Za-z]+-[a-z]+"  # "AI-led", "AI-first"
    r"|[\w-]+\.(com|io|in|ai|co|org|net|dev)(/.*)?)$",  # web domains
    re.I,
)


def _keywords_from(text: str) -> set[str]:
    """Skill-like keywords: known taxonomy skills plus tech-style tokens (React, AWS, C++, Node.js)."""
    lower = text.lower()
    found = {
        s
        for s in ALL_SKILLS
        if re.search(
            rf"(?<![\w]){re.escape(s if s in _CASE_SENSITIVE_SKILLS else s.lower())}(?![\w])",
            text if s in _CASE_SENSITIVE_SKILLS else lower,
        )
    }
    for tok in re.findall(r"\b[A-Z][A-Za-z0-9+#.\-]{1,24}\b", text):
        t = tok.strip(".-")
        if _NOT_A_SKILL.match(t):
            continue
        if len(t) > 1 and t.lower() not in _STOPWORDS and (any(c.isupper() for c in t[1:]) or re.search(r"[+#.\d]", t)):
            found.add(t)
    # Drop fragments of longer skills that were also found ("BI" inside "Power BI")
    return {k for k in found if not any(k != o and re.search(rf"\b{re.escape(k)}\b", o) for o in found)}


def _norm(k: str) -> str:
    return re.sub(r"[^a-z0-9+#]", "", k.lower())


def analyze_resume(text: str, jd_text: str | None = None, target_roles: list[str] | None = None) -> dict[str, Any]:
    """Score a resume out of 100 across six categories and list concrete, fixable issues."""
    lower = text.lower()
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)
    # Only experience and project lines are bullets; a summary paragraph or a skills list is not
    secs = _split_sections(text)
    body = [
        ln
        for ln in secs.get("experience", []) + secs.get("projects", [])
        if re.match(r"^[\-•\*●▪–>•]", ln) or not re.search(r"\b(?:19|20)\d{2}\b|present", ln, re.I)
    ]
    bullets = _bullets("\n".join(body)) if body else _bullets(text)
    issues: list[dict[str, str]] = []
    categories: list[dict[str, Any]] = []

    def add_issue(severity: str, title: str, detail: str, fix: str, category: str) -> None:
        issues.append({"severity": severity, "title": title, "detail": detail, "fix": fix, "category": category})

    # 1. Contact information (10)
    has_email = bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text))
    has_phone = bool(re.search(r"(\+?\d[\d\s\-()]{8,}\d)", text))
    has_linkedin = "linkedin.com" in lower
    has_location = bool(re.search(r"\b(bengaluru|bangalore|hyderabad|pune|mumbai|delhi|gurugram|gurgaon|noida|chennai|kolkata|ahmedabad|india|remote)\b", lower))
    contact_score = 4 * has_email + 3 * has_phone + 2 * has_linkedin + 1 * has_location
    if not has_email:
        add_issue("high", "No email address found", "Recruiters and ATS forms need an email to contact you.", "Add a professional email at the top.", "contact")
    if not has_phone:
        add_issue("high", "No phone number found", "Most Indian employers call before emailing.", "Add your mobile number with country code (+91).", "contact")
    if not has_linkedin:
        add_issue("low", "No LinkedIn URL", "Many recruiters check LinkedIn before shortlisting.", "Add your LinkedIn profile URL.", "contact")
    if not has_location:
        add_issue("low", "No location", "ATS filters often use city to match location requirements.", "Add your city (e.g. Pune, India).", "contact")
    categories.append({"key": "contact", "label": "Contact details", "score": contact_score, "max": 10})

    # 2. Sections (15)
    present = {k: bool(re.search(p, lower)) for k, p in SECTION_PATTERNS.items()}
    section_score = 3 * present["summary"] + 4 * present["experience"] + 3 * present["education"] + 3 * present["skills"] + 2 * (present["projects"] or present["certifications"])
    for key, label, sev in [
        ("experience", "Experience", "high"),
        ("education", "Education", "medium"),
        ("skills", "Skills", "high"),
        ("summary", "Summary", "medium"),
    ]:
        if not present[key]:
            add_issue(sev, f"Missing “{label}” section", f"ATS parsers look for a clearly labelled {label} heading.", f"Add a section titled “{label}”.", "sections")
    if not present["projects"] and not present["certifications"]:
        add_issue("low", "No projects or certifications", "For freshers and career switchers, projects are often the strongest evidence.", "Add 1–3 projects with what you built and the result.", "sections")
    categories.append({"key": "sections", "label": "Standard sections", "score": section_score, "max": 15})

    # 3. Bullet quality (25)
    weak_bullets: list[dict[str, str]] = []
    if bullets:
        verb_start = [b for b in bullets if b.split()[0].lower().strip(",.") in ACTION_VERBS]
        with_metric = [b for b in bullets if has_metric(b)]
        too_long = [b for b in bullets if len(b.split()) > 35]
        weak = [b for b in bullets if any(w in b.lower() for w in WEAK_PHRASES)]
        verb_ratio = len(verb_start) / len(bullets)
        metric_ratio = len(with_metric) / len(bullets)
        bullet_score = round(10 * verb_ratio + 10 * min(1.0, metric_ratio / 0.5) + 5 * (1 - min(1.0, (len(too_long) + len(weak)) / max(1, len(bullets)))))
        if verb_ratio < 0.6:
            add_issue("medium", "Bullets don't start with action verbs", f"Only {round(verb_ratio * 100)}% of bullets start with a strong verb like Built, Led or Reduced.", "Start each bullet with what you did: “Built…”, “Reduced…”, “Led…”.", "bullets")
        if metric_ratio < 0.4:
            add_issue("high", "Few measurable results", f"Only {round(metric_ratio * 100)}% of bullets include a number. Recruiters and ATS rankers favour quantified impact.", "Add real numbers you can stand behind: users, %, time saved, revenue, team size.", "bullets")
        if weak:
            add_issue("medium", "Weak phrasing", f"{len(weak)} bullet(s) use phrases like “responsible for” or “worked on”.", "Replace them with what you actually achieved.", "bullets")
        if too_long:
            add_issue("low", "Very long bullets", f"{len(too_long)} bullet(s) are over 35 words.", "Keep bullets to 1–2 lines (15–30 words).", "bullets")
        for b in bullets:
            reasons = []
            if b.split()[0].lower().strip(",.") not in ACTION_VERBS:
                reasons.append("doesn't start with an action verb")
            if not has_metric(b):
                reasons.append("has no measurable result")
            if any(w in b.lower() for w in WEAK_PHRASES):
                reasons.append("uses weak phrasing")
            if len(reasons) >= 2 or any("weak" in r for r in reasons):
                weak_bullets.append({"text": b, "reason": "; ".join(reasons)})
    else:
        bullet_score = 0
        add_issue("high", "No achievement bullets found", "ATS and recruiters scan bullet points under each role.", "List 3–5 bullets per role describing what you did and the result.", "bullets")
    categories.append({"key": "bullets", "label": "Achievement bullets", "score": bullet_score, "max": 25})

    # 4. Keywords (25)
    resume_kw = _keywords_from(text)
    resume_norm = {_norm(k) for k in resume_kw}
    # Known skills plus tech-style tokens (C++, Node.js, Neo4J); other capitalised words are usually
    # names of brands, colleges or events (InMobi, JNTU, IPL) rather than skills
    known = {sk.lower() for sk in ALL_SKILLS} | _TECH_ACRONYMS

    def is_skill(k: str) -> bool:
        return k.lower() in known or bool(re.search(r"[+#.\d]", k))

    if jd_text and jd_text.strip():
        jd_kw = {k for k in _keywords_from(jd_text) if is_skill(k)}
        matched = sorted(k for k in jd_kw if _norm(k) in resume_norm or k.lower() in lower)
        missing = sorted(k for k in jd_kw if k not in matched)
        coverage = len(matched) / max(1, len(jd_kw))
        keyword_score = round(25 * min(1.0, coverage / 0.75))
        keyword_source = "job_description"
        if missing:
            add_issue("high" if coverage < 0.5 else "medium", "Missing job keywords", f"The job mentions {len(missing)} skills your resume doesn't: {', '.join(missing[:8])}{'…' if len(missing) > 8 else ''}.", "Add the ones you genuinely have, in your skills and in the bullet where you used them. Don't add skills you don't have.", "keywords")
    else:
        # The skills the person listed themselves, plus recognised tools mentioned anywhere
        listed = _parse_skills(secs.get("skills", []), "") if secs.get("skills") else []
        general = {k for k in resume_kw if is_skill(k)} | set(listed)
        matched = sorted(general, key=str.lower)
        missing = []
        keyword_score = round(25 * min(1.0, len(general) / 12))
        keyword_source = "general"
        if len(general) < 8:
            add_issue("medium", "Few recognisable skills", f"We found only {len(general)} skills or tools.", "List the specific tools, languages and platforms you use (e.g. Python, SQL, Figma, AWS).", "keywords")
    categories.append({"key": "keywords", "label": "Keyword match", "score": keyword_score, "max": 25, "source": keyword_source})

    # 5. Formatting and parseability (15)
    fmt = 15
    if word_count < 250:
        fmt -= 5
        add_issue("medium", "Resume is too short", f"About {word_count} words. ATS rankers have little to match against.", "Aim for 400–800 words: more detail on your impact and skills.", "formatting")
    elif word_count > 1100:
        fmt -= 4
        add_issue("medium", "Resume is too long", f"About {word_count} words.", "Keep it to 1 page (freshers) or 2 pages (experienced). Cut older or less relevant detail.", "formatting")
    pronouns = len(re.findall(r"\b(i|me|my|myself)\b", lower))
    if pronouns > 3:
        fmt -= 3
        add_issue("low", "First-person pronouns", f"“I/me/my” appears {pronouns} times.", "Write bullets without pronouns: “Built X”, not “I built X”.", "formatting")
    date_formats = {
        "month_year": len(re.findall(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+\d{4}\b", lower)),
        # "01/2023", but not a grade like "8.6/10"
        "numeric": len(re.findall(r"(?<![\d.])(?:0?[1-9]|1[0-2])/(?:19|20)\d{2}\b", text)),
    }
    if date_formats["month_year"] and date_formats["numeric"]:
        fmt -= 2
        add_issue("low", "Inconsistent date formats", "Mixing “Jan 2023” and “01/2023” can confuse parsers.", "Use one format everywhere, e.g. “Jan 2023 – Mar 2024”.", "formatting")
    odd_chars = len(re.findall(r"[^\x00-\x7F₹–—’‘“”•]", text))
    if odd_chars > 20:
        fmt -= 3
        add_issue("medium", "Special characters or icons", "Icons, emoji and decorative symbols often turn into garbage in ATS parsers.", "Remove icons and use plain text labels (Email:, Phone:).", "formatting")
    if "|" in text and text.count("|") > 15:
        fmt -= 2
        add_issue("low", "Possible tables or columns", "Multi-column layouts and tables are read in the wrong order by many ATS.", "Use a single-column layout. Download our ATS template.", "formatting")
    categories.append({"key": "formatting", "label": "Formatting & parsing", "score": max(0, fmt), "max": 15})

    # 6. Readability (10)
    cliche_hits = [c for c in CLICHES if c in lower]
    sentences = [s for s in re.split(r"[.!?\n]", text) if len(s.split()) > 3]
    avg_len = sum(len(s.split()) for s in sentences) / max(1, len(sentences))
    read = 10 - min(4, 2 * len(cliche_hits)) - (3 if avg_len > 30 else 0)
    if cliche_hits:
        add_issue("low", "Clichés", f"Found: {', '.join(cliche_hits[:5])}.", "Replace adjectives with evidence. Show teamwork through a result instead of saying “team player”.", "readability")
    categories.append({"key": "readability", "label": "Readability", "score": max(0, read), "max": 10})

    total = sum(c["score"] for c in categories)
    severity_order = {"high": 0, "medium": 1, "low": 2}
    issues.sort(key=lambda i: severity_order[i["severity"]])

    return {
        "score": total,
        "grade": "excellent" if total >= 85 else "good" if total >= 70 else "fair" if total >= 50 else "needs_work",
        "categories": categories,
        "issues": issues,
        "keywords": {"matched": matched[:60], "missing": missing[:40], "source": keyword_source},
        "weak_bullets": weak_bullets[:15],
        "stats": {
            "word_count": word_count,
            "bullet_count": len(bullets),
            "sections_found": [k for k, v in present.items() if v],
        },
        "target_roles": target_roles or [],
    }


def _numbers(s: str) -> set[str]:
    return set(re.findall(r"\d+(?:\.\d+)?", s))


def verify_rewrite(original: str, rewritten: str, known_skills: list[str]) -> str | None:
    """Return a reason if the rewrite adds facts that aren't in the original or the profile."""
    new_numbers = _numbers(rewritten) - _numbers(original)
    if new_numbers:
        return f"adds numbers not in your original ({', '.join(sorted(new_numbers))})"
    known = {_norm(k) for k in known_skills} | {_norm(k) for k in _keywords_from(original)}
    new_tools = [k for k in _keywords_from(rewritten) if _norm(k) not in known]
    if new_tools:
        return f"mentions tools not in your profile ({', '.join(sorted(new_tools)[:4])})"
    return None


def profile_snapshot_to_text(snapshot: dict[str, Any]) -> str:
    """Plain-text resume rendering of a profile snapshot, used to scan the master profile."""
    p = snapshot.get("profile", {})
    lines = [
        p.get("full_name") or "",
        p.get("headline") or "",
        " | ".join(x for x in [p.get("email"), p.get("phone"), p.get("city"), *((p.get("links") or {}).values())] if x),
        "Summary",
        p.get("summary") or "",
        "Skills",
        ", ".join(s.get("name", "") for s in snapshot.get("skills", [])),
        "Experience",
    ]
    for e in snapshot.get("experiences", []):
        lines.append(f"{e.get('title', '')} - {e.get('company', '')} ({e.get('start_date') or ''} - {'Present' if e.get('is_current') else e.get('end_date') or ''})")
        for b in e.get("bullets", []):
            lines.append("- " + (b.get("text", "") if isinstance(b, dict) else str(b)))
    if snapshot.get("projects"):
        lines.append("Projects")
        for pr in snapshot["projects"]:
            lines.append(f"{pr.get('title', '')}: {pr.get('description') or ''}")
    lines.append("Education")
    for ed in snapshot.get("educations", []):
        lines.append(f"{ed.get('degree', '')} {ed.get('field', '')} - {ed.get('institution', '')} {ed.get('grade') or ''}")
    if snapshot.get("certifications"):
        lines.append("Certifications")
        for c in snapshot["certifications"]:
            lines.append(c.get("name", ""))
    return "\n".join(x for x in lines if x is not None)


_HEADING_RE = re.compile(
    r"^(?:professional|work|technical|key|core|academic|relevant|personal|career|other)?\s*"
    r"(summary|profile|objective|about me|experience|employment|work history|internships?|"
    r"education|academics?|qualifications?|skills|tech stack|technologies|projects?|"
    r"certifications?|certificates?|courses|achievements|awards)"
    r"(?:\s*(?:&|and)\s*\w+)?\s*:?$",
    re.I,
)


def _split_sections(text: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"_header": []}
    current = "_header"
    for ln in _lines(text):
        m = _HEADING_RE.match(ln)
        if m:
            h = m.group(1).lower()
            current = next((k for k, p in SECTION_PATTERNS.items() if re.search(p, h)), h)
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(ln)
    return sections


_CITIES = (
    "Bengaluru|Bangalore|Hyderabad|Pune|Mumbai|Navi Mumbai|Thane|New Delhi|Delhi|Gurugram|Gurgaon|Noida|"
    "Chennai|Kolkata|Ahmedabad|Jaipur|Kochi|Coimbatore|Indore|Chandigarh|Lucknow|Bhubaneswar|Mysuru|"
    "Mysore|Vadodara|Nagpur|Visakhapatnam|Trivandrum|Thiruvananthapuram|Remote"
)
_CITY_RE = re.compile(rf"\b({_CITIES})\b", re.I)
_LINK_RE = re.compile(r"linkedin\.com|github\.com|https?://|www\.", re.I)
_DATE_RE = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+\d{4}|\d{4}|present|current"
_DEGREE_RE = re.compile(
    r"^(B\.?\s?Tech|B\.?\s?E\.?|M\.?\s?Tech|M\.?\s?E\.?|B\.?\s?Sc|M\.?\s?Sc|BCA|MCA|MBA|PGDM|B\.?\s?Com|M\.?\s?Com|BBA|"
    r"B\.?\s?A\.?|M\.?\s?A\.?|Ph\.?\s?D|Diploma|Bachelor(?:'s)?(?: of [A-Za-z ]+)?|Master(?:'s)?(?: of [A-Za-z ]+)?|"
    r"Class (?:X|XII|10|12)|HSC|SSC)\b\.?(?:\s+(?:in|of)\s+(.+))?",
    re.I,
)
_GRADE_RE = re.compile(
    r"((?:CGPA|GPA|SGPA)\s*[:\-]?\s*\d+(?:\.\d+)?(?:\s*/\s*\d+)?|\d+(?:\.\d+)?\s*(?:/\s*10\s*)?(?:CGPA|GPA)|\d{2}(?:\.\d+)?\s*%)",
    re.I,
)
_ISSUERS = (
    "Google", "Microsoft", "AWS", "Amazon", "Coursera", "Udemy", "Oracle", "IBM", "Meta", "NPTEL", "Cisco",
    "Salesforce", "HubSpot", "Scrum Alliance", "PMI", "Tableau", "Databricks", "Snowflake", "edX", "Infosys",
)


def _strip_location(company: str) -> tuple[str, str]:
    """'Swiggy, Hyderabad' -> ('Swiggy', 'Hyderabad')."""
    parts = [p.strip() for p in re.split(r"\s*[,|]\s*", company) if p.strip()]
    if len(parts) > 1 and _CITY_RE.fullmatch(parts[-1]):
        return ", ".join(parts[:-1]), parts[-1]
    return company, ""


def _parse_education(lines: list[str]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for ln in lines[:6]:
        grade = _GRADE_RE.search(ln)
        years = re.findall(r"\b(?:19|20)\d{2}\b", ln)
        rest = _GRADE_RE.sub("", ln)
        rest = re.sub(r"\(?\b(?:19|20)\d{2}\b\s*(?:[-–]\s*(?:\b(?:19|20)\d{2}\b|present))?\)?", "", rest, flags=re.I)
        parts = [p.strip(" -–|()") for p in re.split(r"\s*[,|]\s*|\s+-\s+", rest) if p.strip(" -–|()")]
        degree, field, institution_parts = "", "", []
        for part in parts:
            m = _DEGREE_RE.match(part)
            if m and not degree:
                degree = part[: m.start(2)].strip().removesuffix(" in").removesuffix(" of") if m.group(2) else part
                field = (m.group(2) or "").strip()
            else:
                institution_parts.append(part)
        entry = {
            "institution": ", ".join(institution_parts),
            "degree": degree,
            "field": field,
            "start_date": years[0] if len(years) > 1 else "",
            "end_date": years[-1] if years else "",
            "grade": grade.group(1).strip() if grade else "",
        }
        # A line with only an institution or only a grade continues the previous entry
        if out and not degree and (not out[-1]["institution"] or not entry["institution"]):
            prev = out[-1]
            prev["institution"] = prev["institution"] or entry["institution"]
            prev["grade"] = prev["grade"] or entry["grade"]
            prev["end_date"] = prev["end_date"] or entry["end_date"]
            continue
        if degree or entry["institution"]:
            out.append(entry)
    return out


def _parse_projects(lines: list[str]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for ln in lines:
        is_bullet = bool(re.match(r"^[\-•\*●▪>]", ln))
        txt = re.sub(r"^[\-•\*●▪>]+\s*", "", ln)
        if is_bullet and out:
            out[-1]["description"] = (out[-1]["description"] + " " + txt).strip()
            continue
        title, _, desc = txt.partition(":") if ":" in txt[:80] else txt.partition(" - ")
        out.append({"title": title.strip()[:128], "description": desc.strip(), "role": "", "url": None, "skills": []})
    for p in out:
        p["skills"] = sorted(_keywords_from(p["description"]))[:8]
    return [p for p in out if p["title"]][:8]


def _parse_certifications(lines: list[str]) -> list[dict[str, Any]]:
    out = []
    for ln in lines[:10]:
        txt = re.sub(r"^[\-•\*●▪>]+\s*", "", ln)
        year = re.search(r"\b(?:19|20)\d{2}\b", txt)
        name = re.sub(r"\s*[(\[,–-]?\s*\b(?:19|20)\d{2}\b\s*[)\]]?\s*$", "", txt).strip(" ,-–")
        issuer = next((i for i in _ISSUERS if re.search(rf"\b{re.escape(i)}\b", name, re.I)), "")
        if name:
            out.append({"name": name[:255], "issuer": issuer, "issue_date": year.group(0) if year else "", "credential_url": None})
    return out


def _parse_skills(lines: list[str], text: str) -> list[str]:
    if lines:
        items: list[str] = []
        for ln in lines:
            ln = re.sub(r"^[\-•\*●▪>]+\s*", "", ln)
            ln = re.sub(r"^[A-Za-z /&]{2,30}:\s*", "", ln)  # "Languages: Python, SQL"
            items += [x.strip(" .") for x in re.split(r"\s*[,|;•]\s*(?![^()]*\))", ln) if x.strip(" .")]
        seen: dict[str, str] = {}
        for it in items:
            if 1 <= len(it) <= 40 and it.lower() not in seen:
                seen[it.lower()] = it
        if seen:
            return list(seen.values())[:40]
    return sorted(_keywords_from(text))[:40]


def heuristic_parse(text: str) -> dict[str, Any]:
    """Rule-based resume parser used when the AI parser is unavailable.

    Less complete than the AI parser, but it only ever returns facts that are in the text.
    """
    sections = _split_sections(text)
    header = sections.get("_header", [])
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    phone = re.search(r"(\+?\d[\d\s\-()]{8,}\d)", text)
    city = _CITY_RE.search(" ".join(header)) or _CITY_RE.search(text)
    name = next((h for h in header if 1 < len(h.split()) <= 5 and "@" not in h and not re.search(r"\d|\|", h)), "")
    if name.isupper():
        name = name.title()
    headline = ""
    for h in header[1:4]:
        seg = h.split("|")[0].strip()
        if seg and seg != name and "@" not in seg and not re.search(r"\d", seg) and not _LINK_RE.search(seg) and not _CITY_RE.fullmatch(seg):
            headline = seg
            break

    experiences: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for ln in sections.get("experience", []):
        is_bullet = bool(re.match(r"^[\-•\*●▪>]", ln))
        if is_bullet and current is not None:
            txt = re.sub(r"^[\-•\*●▪>]+\s*", "", ln)
            current["bullets"].append({"id": f"b{len(experiences)}_{len(current['bullets']) + 1}", "text": txt})
        elif not is_bullet:
            dates = re.findall(rf"({_DATE_RE})", ln, re.I)
            if current is not None and not dates and current["bullets"] and len(ln.split()) > 6:
                # Wrapped bullet text without a bullet marker
                current["bullets"][-1]["text"] += " " + ln
                continue
            head = re.sub(rf"\(?({_DATE_RE})\)?", "", ln, flags=re.I)
            head = re.sub(r"[\s\-–|,]+$", "", head).strip()
            parts = re.split(r"\s+(?:at|@)\s+|\s+[-–|]\s+|\s*,\s*", head, maxsplit=1)
            company, location = _strip_location(parts[1].strip() if len(parts) > 1 else "")
            current = {
                "title": parts[0].strip() if parts else head,
                "company": company,
                "location": location,
                "start_date": dates[0] if dates else "",
                "end_date": dates[1] if len(dates) > 1 and dates[1].lower() not in ("present", "current") else None,
                "is_current": any(d.lower() in ("present", "current") for d in dates),
                "bullets": [],
            }
            experiences.append(current)

    links: dict[str, str] = {}
    for key, pattern in (
        ("linkedin", r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-%]+"),
        ("github", r"(?:https?://)?(?:www\.)?github\.com/[\w\-]+"),
    ):
        m = re.search(pattern, text, re.I)
        if m:
            links[key] = m.group(0) if m.group(0).startswith("http") else f"https://{m.group(0)}"

    experiences = [e for e in experiences if e["title"]]
    return {
        "full_name": name or "Your name",
        "email": email.group(0) if email else "",
        "phone": phone.group(1).strip() if phone else "",
        "city": city.group(1).title() if city else "",
        "state": "",
        "country": "India",
        "headline": headline or (experiences[0]["title"] if experiences else ""),
        "summary": " ".join(sections.get("summary", []))[:1200],
        "total_experience_months": 0,
        "current_title": experiences[0]["title"] if experiences else "",
        "current_company": experiences[0]["company"] if experiences else "",
        "links": links,
        "experiences": experiences,
        "educations": _parse_education(sections.get("education", [])),
        "projects": _parse_projects(sections.get("projects", [])),
        "certifications": _parse_certifications(sections.get("certifications", [])),
        "skills": [
            {"name": s, "normalized_name": s.lower(), "category": "technical"}
            for s in _parse_skills(sections.get("skills", []), text)
        ],
    }
