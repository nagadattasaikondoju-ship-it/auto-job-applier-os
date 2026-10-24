"""Tests for the ATS scanner, rule-based resume parser, document exports and follow-up drafts."""

import io
from datetime import UTC, datetime

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from pypdf import PdfReader

from apps.api.main import app
from packages.core.services.ats import (
    UnsupportedFileError,
    analyze_resume,
    extract_text,
    heuristic_parse,
    verify_rewrite,
)
from packages.core.services.followups import template_follow_up
from packages.documents.export import (
    build_resume_document,
    cover_letter_pdf_bytes,
    resume_docx_bytes,
    resume_pdf_bytes,
)

WEAK = """Priya Sharma
priya@example.com
Experience
Frontend Developer at Acme
- Responsible for working on the website
- Helped with various tasks in the team
Education
B.Tech Computer Science, COEP
"""

STRONG = (
    """Priya Sharma | priya@example.com | +91 98765 43210 | Pune, India | linkedin.com/in/priya
Summary
Frontend engineer with 3 years building React and TypeScript apps for 200k users.
Skills
React, TypeScript, Next.js, Node.js, AWS, Jest, Figma, SQL, Git, Docker, GraphQL, Tailwind
Experience
Frontend Engineer, Acme (Jan 2022 - Present)
- Built a design system in React used by 6 product teams, cutting UI build time by 30%
- Reduced bundle size by 40% with code splitting, improving LCP from 4.1s to 2.3s
- Led migration of 120 pages from JavaScript to TypeScript with zero production incidents
Projects
Expense tracker: Next.js app with 2k monthly users
Education
B.Tech Computer Science, COEP Pune (2021) - 8.4 CGPA
"""
    + " Additional context about scaling frontend systems." * 25
)

SNAPSHOT = {
    "profile": {"full_name": "Priya Sharma", "email": "p@x.com", "city": "Pune", "summary": "Frontend engineer – ₹ “quotes”"},
    "experiences": [
        {"id": "e1", "title": "Frontend Engineer", "company": "Acme", "start_date": "Jan 2022", "is_current": True,
         "bullets": [{"id": "b1", "text": "Built a design system used by 6 teams"}]}
    ],
    "educations": [{"degree": "B.Tech", "field": "CS", "institution": "COEP", "grade": "8.4"}],
    "skills": [{"name": "React"}, {"name": "TypeScript"}],
    "projects": [],
    "certifications": [],
}


def test_weak_resume_scores_low_with_actionable_issues() -> None:
    r = analyze_resume(WEAK, jd_text="Frontend Engineer with React, TypeScript and AWS")
    assert r["score"] < 50
    assert r["grade"] == "needs_work"
    titles = {i["title"] for i in r["issues"]}
    assert "No phone number found" in titles
    assert "Weak phrasing" in titles
    assert {"React", "TypeScript", "AWS"} <= set(r["keywords"]["missing"])
    assert r["issues"][0]["severity"] == "high"  # sorted most severe first


def test_strong_resume_scores_high() -> None:
    r = analyze_resume(STRONG, jd_text="We need React, TypeScript, GraphQL and AWS.")
    assert r["score"] >= 85
    assert r["keywords"]["missing"] == []
    assert sum(c["max"] for c in r["categories"]) == 100


def test_rewrite_verification_blocks_invented_facts() -> None:
    assert verify_rewrite("Helped with the website", "Built the website, improving speed by 40%", []) is not None
    assert "tools" in (verify_rewrite("Worked on the dashboard", "Built the dashboard in GraphQL", []) or "")
    assert verify_rewrite("Worked on React dashboard", "Built a React dashboard for internal teams", ["React"]) is None


def test_heuristic_parse_reads_real_text() -> None:
    parsed = heuristic_parse(
        "Priya Sharma\npriya@example.com | +91 98765 43210 | Pune\nExperience\n"
        "Frontend Engineer at Acme (Jan 2022 - Present)\n- Built a design system in React\nSkills\nReact, AWS\n"
    )
    assert parsed["full_name"] == "Priya Sharma"
    assert parsed["email"] == "priya@example.com"
    assert parsed["city"] == "Pune"
    exp = parsed["experiences"][0]
    assert (exp["title"], exp["company"], exp["is_current"]) == ("Frontend Engineer", "Acme", True)
    assert exp["bullets"][0]["text"] == "Built a design system in React"
    assert {"React", "AWS"} <= {s["name"] for s in parsed["skills"]}


def test_extract_text_formats() -> None:
    d = Document()
    d.add_paragraph("Priya Sharma " * 20)
    buf = io.BytesIO()
    d.save(buf)
    assert "Priya Sharma" in extract_text("cv.docx", buf.getvalue())
    with pytest.raises(UnsupportedFileError):
        extract_text("cv.png", b"x" * 500)
    with pytest.raises(UnsupportedFileError):
        extract_text("cv.txt", b"too short")


def test_pdf_and_docx_exports_contain_resume_text() -> None:
    doc = build_resume_document(SNAPSHOT)
    pdf_text = PdfReader(io.BytesIO(resume_pdf_bytes(doc))).pages[0].extract_text()
    assert "Priya Sharma" in pdf_text
    assert "Built a design system used by 6 teams" in pdf_text
    assert "EXPERIENCE" in pdf_text
    docx = Document(io.BytesIO(resume_docx_bytes(doc)))
    assert any("Frontend Engineer" in p.text for p in docx.paragraphs)
    assert cover_letter_pdf_bytes("Priya", ["p@x.com"], "Dear team,\n\nHello.").startswith(b"%PDF")


def test_tailored_content_overrides_master_bullets() -> None:
    tailored = {
        "summary": "Tailored summary",
        "skills_ordered": ["TypeScript", "React"],
        "experiences": [{"experience_id": "e1", "company": "Acme", "title": "Frontend Engineer",
                         "bullets": [{"source_bullet_id": "b1", "text": "Built a React design system for 6 teams"}]}],
    }
    doc = build_resume_document(SNAPSHOT, tailored)
    assert doc["summary"] == "Tailored summary"
    assert doc["skills"] == ["TypeScript", "React"]
    assert doc["experiences"][0]["bullets"] == ["Built a React design system for 6 teams"]
    assert doc["experiences"][0]["dates"] == "Jan 2022 - Present"


def test_follow_up_template() -> None:
    draft = template_follow_up("Priya Sharma", "Frontend Engineer", "Razorfin", datetime(2026, 10, 1, tzinfo=UTC), "Built a design system")
    assert "Frontend Engineer" in draft["subject"]
    assert "1 October" in draft["body"] and "Razorfin" in draft["body"]
    assert draft["body"].rstrip().endswith("Priya Sharma")


@pytest.mark.asyncio
async def test_resume_upload_rejects_unsupported_files() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        resp = await ac.post(
            "/v1/resumes/upload",
            headers={"X-Dev-User-Id": "dev-user-1"},
            files={"file": ("photo.png", b"\x89PNG" + b"0" * 400, "image/png")},
        )
    assert resp.status_code == 422
    assert resp.json()["error"]["code"] == "UNSUPPORTED_RESUME"
